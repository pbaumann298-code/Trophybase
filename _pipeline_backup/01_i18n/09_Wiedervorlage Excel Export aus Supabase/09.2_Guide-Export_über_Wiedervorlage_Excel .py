#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
09.2_Guide-Export_über_Wiedervorlage_Excel .py (Synchronisiert auf Schema 10/2026 & 10.1 Upload)
- Liest die Game-UUIDs aus game_wiedervorlage_export.xlsx (Schritt 1: 09.1).
- Paginierter Abruf (>1.000 Zeilen) für game_guides & numerische Sortierung für Trophäen.
- Fallback auf game_guides_rohmaterial / staging_game_guides, falls game_guides noch 0 Zeilen hat.
- Trennt Bosse (sheet_type 3 / boss_group) sauber aus game_guides in Sheet 3 auf (game_bosses existiert nicht mehr).
- Überschreibt NICHT mehr die DLC-Namen in game_achievements.spielname.
- Kein falscher DE->EN Fallback bei guide_tip / Guide-Tip-lang (verhindert deutsche Tipps im EN-Key beim Re-Upload).
"""

import ast
import json
import os
import re
import sys
from typing import Any, Dict, List
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv
from supabase import Client, create_client

# --- SETUP ---
load_dotenv()
SUPABASE_URL = os.getenv("SUPABASE_URL") or os.getenv("VITE_SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:
    sys.exit(
        "❌ Fehler: SUPABASE_URL oder SUPABASE_SERVICE_ROLE_KEY fehlt in der .env!"
    )

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

_BASE_DIR = os.path.dirname(os.path.abspath(__file__))
INPUT_EXCEL = os.path.join(_BASE_DIR, "game_wiedervorlage_export.xlsx")
OUTPUT_DIR = os.path.join(_BASE_DIR, "exported_from_db")

os.makedirs(OUTPUT_DIR, exist_ok=True)

# --- KONFIGURATION SPALTEN (100% abgestimmt auf 10.1_Upload_Guide_i18n.py) ---
COLS_SHEET_0 = [
    "game_id",
    "platform_game_id",
    "platform_achievement_id",
    "trophy_id",
    "trophy_name",
    "trophy_desc",
    "trophy_type",
    "trophy_gruppe",
    "guide_tip",
    "Guide-Tip-lang",
    "video_url",
    "timestamp",
    "spielname",
]

COLS_SHEET_1_2 = [
    "platform_game_id",
    "game_id",
    "guide_id",
    "localisation",
    "chronological_group",
    "category_group",
    "boss_group",
    "video_chapter",
    "item_name",
    "timestamp",
    "video_url",
    "prod_note",
    "is_trophy_relevant",
    "trophy_id",
]

COLS_SHEET_3 = [
    "platform_game_id",
    "game_id",
    "boss_id",
    "localisation",
    "chronological_group",
    "category_group",
    "boss_group",
    "video_chapter",
    "boss_name",
    "timestamp",
    "video_url",
    "prod_note",
    "is_trophy_relevant",
    "trophy_id",
]


# --- HILFSFUNKTIONEN ---
def clean_npwr_string(raw_id: Any) -> str:
    if not raw_id:
        return ""
    if isinstance(raw_id, list):
        return str(raw_id[0]).strip("[]'\" ") if raw_id else ""
    val = str(raw_id).strip("[]'\" ")
    if "," in val:
        val = val.split(",")[0].strip("[]'\" ")
    return val


def clean_yt_url(url: Any) -> str:
    """Entfernt Timestamp-Parameter (?t=..., &t=..., ?start=...) aus der URL."""
    if not url or pd.isna(url):
        return ""
    val = str(url).strip()
    if not val or val.lower() in ("nan", "none"):
        return ""

    cleaned = re.sub(r"[?&](?:t|start)=\d+s?", "", val)
    cleaned = re.sub(r"\?&+", "?", cleaned)
    cleaned = re.sub(r"[?&]+$", "", cleaned)

    if "?" not in cleaned and "&" in cleaned:
        cleaned = cleaned.replace("&", "?", 1)

    return cleaned


def parse_jsonb_value(val: Any) -> Any:
    if val is None or val == "":
        return ""
    if isinstance(val, (dict, list)):
        return val
    if isinstance(val, str):
        cleaned = val.strip()
        if (cleaned.startswith("{") and cleaned.endswith("}")) or (
            cleaned.startswith("[") and cleaned.endswith("]")
        ):
            try:
                return json.loads(cleaned)
            except Exception:
                pass
            try:
                return ast.literal_eval(cleaned)
            except Exception:
                return val
    return val


def extract_localized_text(val: Any, lang: str = "de", exact: bool = False) -> Any:
    """
    Löst JSONB-Sprachmaps auf.
    Bei exact=True (z. B. für guide_tip) wird KEIN Fallback auf eine andere Sprache genutzt,
    damit beim EN-Export kein deutscher Tipp im englischen Feld landet!
    """
    if isinstance(val, (bool, int, float)):
        return val

    parsed = parse_jsonb_value(val)
    if parsed is None or parsed == "":
        return ""

    if isinstance(parsed, dict):
        if lang in parsed and parsed[lang] not in (None, ""):
            return str(parsed[lang]).strip()
        if exact:
            return ""
        fallback_lang = "en" if lang == "de" else "de"
        if fallback_lang in parsed and parsed[fallback_lang] not in (None, ""):
            return str(parsed[fallback_lang]).strip()
        first_val = next((v for v in parsed.values() if v not in (None, "")), "")
        return str(first_val).strip()

    if isinstance(parsed, list):
        return ", ".join(
            str(extract_localized_text(x, lang=lang, exact=exact)) for x in parsed
        )

    # Wenn es ein flacher String ist (z. B. aus staging_game_guides), bei exact=True nur für DE ausgeben
    if exact and lang != "de":
        return ""
    return str(parsed).strip()


def resolve_sheet_types(raw_sheet_type: Any, row: dict) -> List[int]:
    parsed = parse_jsonb_value(raw_sheet_type)
    types = []
    if isinstance(parsed, list):
        for x in parsed:
            try:
                types.append(int(x))
            except (ValueError, TypeError):
                pass
    elif isinstance(parsed, (int, float)):
        types.append(int(parsed))

    if not types:
        if extract_localized_text(row.get("chronological_group")):
            types.append(1)
        if extract_localized_text(row.get("category_group")):
            types.append(2)
        if extract_localized_text(row.get("boss_group")):
            types.append(3)
    return sorted(set(types)) if types else [1]


def has_language_content(val: Any, lang: str = "en") -> bool:
    parsed = parse_jsonb_value(val)
    if isinstance(parsed, dict):
        return bool(str(parsed.get(lang) or "").strip())
    return False


def fetch_all_rows_paginated(table_name: str, game_uuid: str) -> List[Dict]:
    all_rows = []
    page_size = 1000
    offset = 0
    while True:
        res = (
            supabase.table(table_name)
            .select("*")
            .eq("game_id", game_uuid)
            .range(offset, offset + page_size - 1)
            .execute()
        )
        data = res.data or []
        all_rows.extend(data)
        if len(data) < page_size:
            break
        offset += page_size
    return all_rows


def sort_number(val: Any) -> int:
    m = re.search(r"\d+", str(val or ""))
    return int(m.group(0)) if m else 999999


def load_uuids_from_excel(file_path: str) -> List[str]:
    if not os.path.isfile(file_path):
        sys.exit(
            f"❌ Eingabedatei '{os.path.basename(file_path)}' wurde im Verzeichnis nicht gefunden!"
        )

    try:
        df_target = pd.read_excel(file_path)
    except Exception as e:
        sys.exit(f"❌ Fehler beim Lesen der Excel-Datei '{file_path}': {e}")

    if df_target.empty:
        sys.exit(f"❌ Die Datei '{os.path.basename(file_path)}' enthält keine Daten.")

    target_col = None
    for candidate in ["id", "game_id", "uuid", "spiele_id"]:
        for col in df_target.columns:
            if str(col).strip().lower() == candidate:
                target_col = col
                break
        if target_col:
            break

    if not target_col:
        target_col = df_target.columns[0]
        print(
            f"ℹ️ Keine explizite 'ID'-Spalte gefunden. Nutze erste Spalte: '{target_col}'"
        )

    uuids = []
    uuid_pattern = re.compile(r"^[0-9a-fA-F-]{36}$")
    for raw_val in df_target[target_col].dropna():
        val = str(raw_val).strip()
        if uuid_pattern.match(val):
            uuids.append(val)
        else:
            print(f"⚠️ Überspringe ungültigen Eintrag in Spalte '{target_col}': {val}")

    if not uuids:
        sys.exit(
            f"❌ Keine gültigen UUIDs in '{os.path.basename(file_path)}' gefunden."
        )

    return list(dict.fromkeys(uuids))


def ensure_columns(df: pd.DataFrame, expected_cols: List[str]) -> pd.DataFrame:
    for col in expected_cols:
        if col not in df.columns:
            df[col] = ""
    return df[expected_cols]


def style_excel_file(file_path: str):
    if not os.path.exists(file_path):
        return
    wb = load_workbook(file_path)
    font_family = "Segoe UI"
    header_font = Font(name=font_family, size=11, bold=True, color="FFFFFF")
    data_font = Font(name=font_family, size=10, color="333333")
    thin_side = Side(border_style="thin", color="D3D3D3")
    data_border = Border(
        left=thin_side, right=thin_side, top=thin_side, bottom=thin_side
    )

    fills = {
        "sheet_name=0": PatternFill(
            start_color="1B365D", end_color="1B365D", fill_type="solid"
        ),
        "sheet_name=1": PatternFill(
            start_color="2E5A44", end_color="2E5A44", fill_type="solid"
        ),
        "sheet_name=2": PatternFill(
            start_color="7A5C13", end_color="7A5C13", fill_type="solid"
        ),
        "sheet_name=3": PatternFill(
            start_color="4A2828", end_color="4A2828", fill_type="solid"
        ),
    }
    zebra_fill = PatternFill(
        start_color="F9FAFC", end_color="F9FAFC", fill_type="solid"
    )

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        ws.views.sheetView[0].showGridLines = True
        header_fill = fills.get(
            sheet_name,
            PatternFill(start_color="4F5D75", end_color="4F5D75", fill_type="solid"),
        )

        if ws.max_row >= 1:
            for cell in ws[1]:
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
            ws.row_dimensions[1].height = 28

        ws.freeze_panes = "A2"

        for row_idx in range(2, ws.max_row + 1):
            is_zebra = row_idx % 2 == 0
            ws.row_dimensions[row_idx].height = 20
            for col_idx in range(1, ws.max_column + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.font = data_font
                cell.border = data_border
                if is_zebra:
                    cell.fill = zebra_fill

                col_name = str(ws.cell(row=1, column=col_idx).value or "").lower()
                if any(x in col_name for x in ["id", "timestamp", "relevant"]):
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif "url" in col_name or "link" in col_name:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
                    if cell.value and str(cell.value).startswith("http"):
                        cell.font = Font(
                            name=font_family,
                            size=10,
                            color="0000EE",
                            underline="single",
                        )
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                if cell.value:
                    val_str = str(cell.value)
                    max_len = max(
                        max_len, 15 if val_str.startswith("http") else len(val_str)
                    )
            ws.column_dimensions[col_letter].width = max(min(max_len + 3, 50), 12)

    wb.save(file_path)
    wb.close()


# --- PROCESSING HELPERS ---
SKIP_COLS = {
    "game_id",
    "platform_game_id",
    "platform_achievement_id",
    "guide_id",
    "trophy_id",
    "boss_id",
    "ist_versteckt",
    "is_trophy_relevant",
    "timestamp",
    "video_url",
    "prod_note",
}

# Spalten, die niemals auf eine andere Sprache zurückfallen dürfen
EXACT_LANG_COLS = {"guide_tip", "Guide-Tip-lang"}


def clean_dataframe_jsonb(df: pd.DataFrame, lang: str) -> pd.DataFrame:
    if df.empty:
        return df
    for col in df.columns:
        if col not in SKIP_COLS:
            is_exact = col in EXACT_LANG_COLS
            df[col] = df[col].apply(
                lambda x: extract_localized_text(x, lang=lang, exact=is_exact)
            )
    return df


def split_guides_and_bosses(
    all_guide_rows: List[Dict],
) -> tuple[List[Dict], List[Dict]]:
    """
    Teilt die Zeilen aus game_guides (bzw. Rohmaterial/Staging) auf:
    - Sheet 1: Zeilen mit sheet_type 1 oder 2 (oder reine Walkthrough/Collectible-Zeilen)
    - Sheet 3: Zeilen mit sheet_type 3 oder gefüllter boss_group (gemappt auf boss_id & boss_name)
    """
    sheet_1_rows = []
    sheet_3_rows = []

    for r in all_guide_rows:
        st = resolve_sheet_types(r.get("sheet_type"), r)
        has_boss = 3 in st or bool(extract_localized_text(r.get("boss_group")))
        has_walk_or_coll = (1 in st) or (2 in st)

        if has_walk_or_coll or not has_boss:
            sheet_1_rows.append(dict(r))

        if has_boss:
            boss_copy = dict(r)
            boss_copy["boss_id"] = r.get("guide_id")
            boss_copy["boss_name"] = r.get("item_name")
            sheet_3_rows.append(boss_copy)

    return sheet_1_rows, sheet_3_rows


def process_data_for_lang(
    lang: str,
    raw_title: Any,
    raw_achievements: List[Dict],
    raw_guides: List[Dict],
    raw_bosses: List[Dict],
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    resolved_title = extract_localized_text(raw_title, lang=lang)

    # 1. Achievements (Sheet 0)
    df_tab_0 = (
        pd.DataFrame(raw_achievements)
        if raw_achievements
        else pd.DataFrame(columns=COLS_SHEET_0)
    )
    if not df_tab_0.empty:
        df_tab_0["trophy_id"] = df_tab_0["platform_achievement_id"]
        df_tab_0 = clean_dataframe_jsonb(df_tab_0, lang=lang)
        # DLC-Gruppennamen in spielname erhalten! Nur wenn leer, auf Spieltitel zurückfallen:
        df_tab_0["spielname"] = df_tab_0["spielname"].apply(
            lambda x: x if str(x or "").strip() else resolved_title
        )
    df_tab_0 = ensure_columns(df_tab_0, COLS_SHEET_0)

    # 2. Guides (Sheet 1 & 2)
    df_guides = (
        pd.DataFrame(raw_guides) if raw_guides else pd.DataFrame(columns=COLS_SHEET_1_2)
    )
    if not df_guides.empty:
        df_guides = clean_dataframe_jsonb(df_guides, lang=lang)
    df_tab_1 = ensure_columns(df_guides.copy(), COLS_SHEET_1_2)
    df_tab_2 = pd.DataFrame(columns=COLS_SHEET_1_2)

    # 3. Bosses (Sheet 3)
    df_bosses = (
        pd.DataFrame(raw_bosses) if raw_bosses else pd.DataFrame(columns=COLS_SHEET_3)
    )
    if not df_bosses.empty:
        df_bosses = clean_dataframe_jsonb(df_bosses, lang=lang)
    df_tab_3 = ensure_columns(df_bosses, COLS_SHEET_3)

    # 4. URLs säubern: Timestamps entfernen, damit nur die blanke Video-URL exportiert wird
    for df in [df_tab_0, df_tab_1, df_tab_2, df_tab_3]:
        if "video_url" in df.columns:
            df["video_url"] = df["video_url"].apply(clean_yt_url)

    return df_tab_0, df_tab_1, df_tab_2, df_tab_3


def save_excel_package(
    file_path: str,
    sheets: tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, pd.DataFrame],
):
    df_0, df_1, df_2, df_3 = sheets
    with pd.ExcelWriter(file_path, engine="openpyxl") as writer:
        df_0.to_excel(writer, sheet_name="sheet_name=0", index=False)
        df_1.to_excel(writer, sheet_name="sheet_name=1", index=False)
        df_2.to_excel(writer, sheet_name="sheet_name=2", index=False)
        df_3.to_excel(writer, sheet_name="sheet_name=3", index=False)
    style_excel_file(file_path)


# --- MAIN LOGIC ---
def run_export():
    uuids = load_uuids_from_excel(INPUT_EXCEL)
    print(
        f"🎯 Geladen: {len(uuids)} Spiele aus '{os.path.basename(INPUT_EXCEL)}'. Starte Export...\n"
    )

    for idx, game_uuid in enumerate(uuids, start=1):
        res_game = (
            supabase.table("games")
            .select("spieltitel, platform_game_id")
            .eq("id", game_uuid)
            .limit(1)
            .execute()
        )
        if not res_game.data:
            print(
                f"❌ {idx}/{len(uuids)}: UUID {game_uuid} nicht in Tabelle 'games' gefunden."
            )
            continue

        game_data = res_game.data[0]
        raw_title = game_data.get("spieltitel")
        npwr_id = clean_npwr_string(game_data.get("platform_game_id"))

        # 1. Trophäen laden & numerisch sortieren (1, 2, ..., 10 statt 1, 10, 2)
        raw_achievements = fetch_all_rows_paginated("game_achievements", game_uuid)
        raw_achievements.sort(
            key=lambda r: sort_number(r.get("platform_achievement_id"))
        )

        # 2. Guides paginiert laden (mit Fallback auf Rohmaterial & Staging, falls noch unveröffentlicht)
        all_guide_rows = fetch_all_rows_paginated("game_guides", game_uuid)
        source_info = "game_guides"

        if not all_guide_rows:
            roh_rows = fetch_all_rows_paginated("game_guides_rohmaterial", game_uuid)
            if roh_rows:
                # Bevorzuge MATCHED-Batch, falls vorhanden
                matched = [
                    r
                    for r in roh_rows
                    if str(r.get("batch_status") or "").upper() == "MATCHED"
                ]
                all_guide_rows = matched if matched else roh_rows
                source_info = "game_guides_rohmaterial"
            else:
                staging_rows = fetch_all_rows_paginated(
                    "staging_game_guides", game_uuid
                )
                if staging_rows:
                    all_guide_rows = staging_rows
                    source_info = "staging_game_guides"

        all_guide_rows.sort(key=lambda r: sort_number(r.get("guide_id")))
        raw_guides, raw_bosses = split_guides_and_bosses(all_guide_rows)

        has_en = has_language_content(raw_title, "en") or any(
            has_language_content(ach.get("trophy_name"), "en")
            or has_language_content(ach.get("trophy_desc"), "en")
            for ach in raw_achievements
        )

        # --- A) DEUTSCHER EXPORT ---
        title_de = extract_localized_text(raw_title, lang="de") or "Unbekanntes_Spiel"
        safe_title_de = re.sub(r'[\\/*?:"<>|]', "", title_de).strip()
        filename_de = (
            f"{npwr_id} - {safe_title_de}.xlsx" if npwr_id else f"{safe_title_de}.xlsx"
        )
        path_de = os.path.join(OUTPUT_DIR, filename_de)

        sheets_de = process_data_for_lang(
            "de", raw_title, raw_achievements, raw_guides, raw_bosses
        )
        save_excel_package(path_de, sheets_de)
        print(
            f"✅ [{idx}/{len(uuids)}] Exportiert (DE): {filename_de} "
            f"({len(raw_guides)} Guide-Zeilen, {len(raw_bosses)} Bosse aus {source_info})"
        )

        # --- B) ENGLISCHER EXPORT (falls vorhanden) ---
        if has_en:
            title_en = extract_localized_text(raw_title, lang="en") or safe_title_de
            safe_title_en = re.sub(r'[\\/*?:"<>|]', "", title_en).strip()
            filename_en = (
                f"{npwr_id} - {safe_title_en}_en.xlsx"
                if npwr_id
                else f"{safe_title_en}_en.xlsx"
            )
            path_en = os.path.join(OUTPUT_DIR, filename_en)

            sheets_en = process_data_for_lang(
                "en", raw_title, raw_achievements, raw_guides, raw_bosses
            )
            save_excel_package(path_en, sheets_en)
            print(f"🌐 [{idx}/{len(uuids)}] Exportiert (EN): {filename_en}")


if __name__ == "__main__":
    run_export()
