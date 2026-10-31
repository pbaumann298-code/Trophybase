import os
import json
import sys
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from supabase import create_client, Client
from dotenv import load_dotenv

load_dotenv()
SUPABASE_URL = os.getenv("SUPABASE_URL")
SUPABASE_KEY = os.getenv("SUPABASE_SERVICE_ROLE_KEY")

if not SUPABASE_URL or not SUPABASE_KEY:
    sys.exit(
        "❌ Fehler: SUPABASE_URL oder SUPABASE_SERVICE_ROLE_KEY fehlt in der .env!"
    )

supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)
TABLE_NAME = "game_wiedervorlage"
OUTPUT_FILE = "game_wiedervorlage_export.xlsx"


def fetch_and_export():
    # 1. Supabase Client initialisieren
    supabase: Client = create_client(SUPABASE_URL, SUPABASE_KEY)

    # 2. Relevante Spalten selektieren
    response = (
        supabase.table(TABLE_NAME)
        .select("id, platform_game_id, spieltitel, freitext")
        .execute()
    )

    data = response.data
    if not data:
        print("Keine Datensätze gefunden.")
        return

    # 3. Daten aufbereiten (JSON / Arrays für Excel lesbar machen)
    processed_rows = []
    for row in data:
        # platform_game_id: Falls JSON-Array oder Liste vorliegt -> komma-separierter String
        raw_platform_id = row.get("platform_game_id")
        if isinstance(raw_platform_id, str):
            try:
                parsed = json.loads(raw_platform_id)
                platform_id_str = (
                    ", ".join(parsed) if isinstance(parsed, list) else raw_platform_id
                )
            except (json.JSONDecodeError, TypeError):
                platform_id_str = raw_platform_id
        elif isinstance(raw_platform_id, list):
            platform_id_str = ", ".join(map(str, raw_platform_id))
        else:
            platform_id_str = str(raw_platform_id or "")

        # spieltitel: Fallback-Hierarchie (de -> en -> Fallback auf String/Rohdaten)
        raw_titel = row.get("spieltitel")
        if isinstance(raw_titel, str):
            try:
                parsed_titel = json.loads(raw_titel)
                if isinstance(parsed_titel, dict):
                    titel_str = (
                        parsed_titel.get("de")
                        or parsed_titel.get("en")
                        or next(iter(parsed_titel.values()), "")
                    )
                else:
                    titel_str = raw_titel
            except (json.JSONDecodeError, TypeError):
                titel_str = raw_titel
        elif isinstance(raw_titel, dict):
            titel_str = (
                raw_titel.get("de")
                or raw_titel.get("en")
                or next(iter(raw_titel.values()), "")
            )
        else:
            titel_str = str(raw_titel or "")

        processed_rows.append(
            {
                "ID": row.get("id"),
                "Platform Game ID": platform_id_str,
                "Spieltitel": titel_str,
                "Freitext": row.get("freitext") or "",
            }
        )

    # 4. DataFrame erstellen
    df = pd.DataFrame(processed_rows)

    # 5. In Excel schreiben & formatieren
    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        sheet_name = "Wiedervorlage"
        df.to_excel(writer, index=False, sheet_name=sheet_name)
        worksheet = writer.sheets[sheet_name]

        # Styling
        header_fill = PatternFill(
            start_color="1F2937", end_color="1F2937", fill_type="solid"
        )  # Slate/Dark Gray
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="E5E7EB"),
            right=Side(style="thin", color="E5E7EB"),
            top=Side(style="thin", color="E5E7EB"),
            bottom=Side(style="thin", color="E5E7EB"),
        )

        # Header formatieren
        for col_idx in range(1, len(df.columns) + 1):
            cell = worksheet.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left", vertical="center")

        # Zeilen formatieren & Spaltenbreiten anpassen
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions

        for row in worksheet.iter_rows(
            min_row=2, max_row=len(df) + 1, min_col=1, max_col=len(df.columns)
        ):
            for cell in row:
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        for col in worksheet.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            worksheet.column_dimensions[col_letter].width = min(
                max(max_len + 4, 14), 50
            )

    print(f"Export erfolgreich abgeschlossen: {OUTPUT_FILE} ({len(df)} Zeilen)")


if __name__ == "__main__":
    fetch_and_export()
