#!/usr/bin/env python3
r"""
08.1_Rohdatenveredler.py (UPGRADE v2)
Läuft automatisch über C:\Temp\TrophyBase\01_i18n\08_Guide_Ordner_in_Bearbeitung:
1. Zieht Kombi-Zeilen mit '-', 'to', 'bis', '&', ',' und 'and' (z.B. 'Document #9 to #15') in Einzelzeilen auseinander.
2. Wandelt '<Untitled Chapter 1>' & 'Start' bei 00:00 in saubere Walkthrough-Startanker um.
3. Leitet automatisch 'chronological_group' (Part/Kapitel) und 'category_group' ab.
4. Erkennt kapitelweise Nummerierungs-Resets (#1..#N pro Part) UND findet fehlende Nummern (wie Document #2),
   fügt sie als gelb markierte Prüf-Zeile im richtigen Zeitfenster ein und formatiert die Excel sauber.
"""

import os
import re
import sys
import glob
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

DEFAULT_TARGET_DIR = r"C:\Temp\TrophyBase\01_i18n\08_Guide_Ordner_in_Bearbeitung"

# Matcht "#9 to #15", "#15-17", "Document (Drawings) #9 to #15"
RANGE_UNIVERSAL_RE = re.compile(
    r"^(.*?#?\s*)(\d+)\s*(?:[-–—]|to|bis)\s*#?\s*(\d+)(.*)$", re.I
)
# Matcht "Document #1 & #2", "Transfer tank #6, #7 & #8"
MULTI_UNIVERSAL_RE = re.compile(
    r"^([A-Za-zÄÖÜäöüß\s\(\)\-]+?)\s*#?\s*(\d+(?:\s*(?:,|&|and|und)\s*#?\s*\d+)+)(.*)$",
    re.I,
)
FRACTION_RE = re.compile(r"\b(\d+)\s*/\s*(\d+)\b")
ITEM_NUM_RE = re.compile(r"^([A-Za-zÄÖÜäöüß\s\(\)\-]+?)\s*#(\d+)\b", re.I)

UNTITLED_NAMES = {
    "<untitled chapter 1>",
    "<untitled 1 chapter>",
    "untitled chapter",
    "start",
    "keine timestamps gefunden",
}

# Automatische Kategorie-Erkennung für englische/deutsche Roh-Imports
AUTO_CATEGORY_RULES = [
    (
        re.compile(
            r"\b(document|note|letter|newspaper|diary|file|audio\s*log|recording)\b",
            re.I,
        ),
        "Dokumente, Notizen & Akten",
    ),
    (
        re.compile(r"\b(toy|figurine|doll|plush|statue|bobblehead)\b", re.I),
        "Spielzeuge & Figuren",
    ),
    (
        re.compile(r"\b(mixtape|cassette|tape|vinyl|record|music|song)\b", re.I),
        "Mixtapes & Musik",
    ),
    (
        re.compile(r"\b(briefcase|safe|chest|lockbox|locker|trinket)\b", re.I),
        "Tresore, Koffer & Truhen",
    ),
    (
        re.compile(r"\b(scrap|loot|gold|coin|money|ammo|resource)\b", re.I),
        "Schrott & Ressourcen",
    ),
    (
        re.compile(
            r"\b(relic|artifact|artefact|antique|totem|crest|emblem|medallion)\b", re.I
        ),
        "Relikte, Artefakte & Embleme",
    ),
    (
        re.compile(r"\b(key\s*item|keycard|key|fuse|valve|crank)\b", re.I),
        "Schlüssel-Gegenstände (Key Items)",
    ),
    (re.compile(r"\b(puzzle|riddle)\b", re.I), "Puzzleteile & Rätsel"),
    (
        re.compile(r"\b(weapon|gun|rifle|pistol|shotgun|sword|melee|knife)\b", re.I),
        "Waffen & Ausrüstung",
    ),
    (
        re.compile(r"\b(outfit|costume|skin|armor|armour|clothes|hat|mask)\b", re.I),
        "Outfits & Rüstungen",
    ),
    (re.compile(r"\b(survivor|hostage|npc|rescue)\b", re.I), "Überlebende & NPCs"),
    (
        re.compile(r"\b(boss|psychopath|top\s*dog|aberration)\b", re.I),
        "Bosse & Spezialgegner",
    ),
    (
        re.compile(r"\b(sticker|photo|picture|camera|poster|graffiti)\b", re.I),
        "Fotos, Sticker & Poster",
    ),
    (
        re.compile(r"\b(achievement|trophy|trophäe)\b", re.I),
        "Spezifische Trophäen-Aktionen",
    ),
]


def extract_clean_chapter_name(video_chapter: str) -> str:
    """Zieht aus langen YouTube-Titeln den sauberen Part-/Kapitel-Namen."""
    s = str(video_chapter or "").strip()
    if not s or s.lower() == "nan":
        return "Hauptspiel"

    # Muster 1: "Part 01 The Ranch" oder "Part 1 - Prologue"
    m_part = re.search(
        r"\b(Part\s*\d+|Chapter\s*\d+|Kapitel\s*\d+|#\d+)\s*[-–—:]?\s*(.*)$", s, re.I
    )
    if m_part:
        tag = m_part.group(1).strip()
        rest = re.sub(
            r"\s*\((?:All\s+collectibles|100%|Walkthrough).*?\)",
            "",
            m_part.group(2),
            flags=re.I,
        ).strip(" -–—:|")
        return f"{tag}: {rest}" if rest else tag

    # Muster 2: "All XYZ Location - ABC Trophy"
    m_all = re.search(r"\bAll\s+([A-Za-z0-9\s]+?)\s+Location", s, re.I)
    if m_all:
        return f"Alle {m_all.group(1).strip()}"

    return s[:60]


def infer_category_group(item_name: str, current_cat: str) -> str:
    cur = str(current_cat or "").strip()
    if cur and cur.lower() not in ("nan", "sammelobjekte", ""):
        return cur
    s = str(item_name or "").strip()
    if s.startswith("🎬"):
        return "Komplett-Walkthrough (Video-Starts ab 00:00)"
    for pattern, label in AUTO_CATEGORY_RULES:
        if pattern.search(s):
            return label
    return cur if cur and cur.lower() != "nan" else "Sammelobjekte"


def expand_and_enrich_rows(df1: pd.DataFrame) -> pd.DataFrame:
    expanded = []
    sub_section_counter = {}
    last_item_nums = {}
    last_vchap = None

    for _, r in df1.iterrows():
        r_dict = r.to_dict()
        raw_iname = str(r_dict.get("item_name") or "").strip()
        vchap = str(r_dict.get("video_chapter") or "").strip()

        # Normiere Abstände vor/nach Raute: "Toy#1" -> "Toy #1", "Document # 10" -> "Document #10"
        iname = re.sub(r"#\s*(\d+)", r"#\1", raw_iname)
        iname = re.sub(r"([A-Za-z\)])#(\d+)", r"\1 #\2", iname)

        base_chap = extract_clean_chapter_name(vchap)
        if vchap != last_vchap:
            last_vchap = vchap
            last_item_nums = {}
            sub_section_counter[base_chap] = 1

        # Prüfe, ob mitten im selben Video die Nummerierung wieder auf #1 zurückspringt (wie in A Quiet Place Part 04!)
        m_num_check = ITEM_NUM_RE.match(iname)
        if m_num_check and iname.lower() not in UNTITLED_NAMES:
            itype_key = m_num_check.group(1).strip().lower()
            cur_num = int(m_num_check.group(2))
            prev_num = last_item_nums.get(itype_key, 0)
            if cur_num == 1 and prev_num >= 2 and len(last_item_nums) >= 2:
                # Neues Unterkapitel im selben Video erkannt!
                sub_section_counter[base_chap] = (
                    sub_section_counter.get(base_chap, 1) + 1
                )
                last_item_nums = {}
            last_item_nums[itype_key] = cur_num

        sec_idx = sub_section_counter.get(base_chap, 1)
        chap_label = f"{base_chap} (Abschnitt {sec_idx})" if sec_idx > 1 else base_chap

        cur_chrono = str(r_dict.get("chronological_group") or "").strip()
        if not cur_chrono or cur_chrono.lower() in ("nan", "hauptspiel"):
            r_dict["chronological_group"] = chap_label

        # 1. 00:00 Start / Untitled Chapter
        if iname.lower() in UNTITLED_NAMES:
            c_row = dict(r_dict)
            c_row["item_name"] = f"🎬 Walkthrough-Start: {chap_label} (ab 00:00)"
            c_row["category_group"] = "Komplett-Walkthrough (Video-Starts ab 00:00)"
            expanded.append(c_row)
            continue

        # 2. Bereiche wie "Document #9 to #15" oder "PP Sticker #15-17"
        m_rng = RANGE_UNIVERSAL_RE.match(iname)
        if m_rng:
            prefix, s_str, e_str, suffix = m_rng.groups()
            s_n, e_n = int(s_str), int(e_str)
            if 1 < (e_n - s_n + 1) <= 35:
                if not prefix.strip().endswith("#"):
                    prefix = prefix.rstrip() + " #"
                for n in range(s_n, e_n + 1):
                    c_row = dict(r_dict)
                    item_txt = (
                        f"{prefix}{n}{suffix} (Sammel-Spot #{s_n}–#{e_n})".strip()
                    )
                    c_row["item_name"] = item_txt
                    c_row["category_group"] = infer_category_group(
                        item_txt, c_row.get("category_group")
                    )
                    expanded.append(c_row)
                    # Update tracker so #1..#9 sets prev_num = 9
                    m_sub = ITEM_NUM_RE.match(f"{prefix}{n}")
                    if m_sub:
                        last_item_nums[m_sub.group(1).strip().lower()] = n
                continue

        # 3. Mehrfach-Nennungen wie "Document #1 & #2" oder "Transfer tank #6, #7 & #8"
        m_mul = MULTI_UNIVERSAL_RE.match(iname)
        if m_mul:
            prefix, nums_blob, suffix = m_mul.groups()
            nums = [int(x) for x in re.findall(r"\d+", nums_blob)]
            if len(nums) > 1:
                for n in nums:
                    c_row = dict(r_dict)
                    item_txt = f"{prefix.strip()} #{n}{suffix} (Kombi-Spot)".strip()
                    c_row["item_name"] = item_txt
                    c_row["category_group"] = infer_category_group(
                        item_txt, c_row.get("category_group")
                    )
                    expanded.append(c_row)
                    last_item_nums[prefix.strip().lower()] = n
                continue

        r_dict["item_name"] = iname
        r_dict["category_group"] = infer_category_group(
            iname, r_dict.get("category_group")
        )
        expanded.append(r_dict)

    return pd.DataFrame(expanded)


def detect_and_insert_gaps(df1: pd.DataFrame, filename: str) -> pd.DataFrame:
    """Prüft pro Abschnitt alle #1..#N Serien auf vergessene Timestamps und fügt Warn-Zeilen ein."""
    if df1.empty:
        return df1

    final_rows = []
    total_gaps = 0

    for chap, grp in df1.groupby("chronological_group", sort=False):
        grp_rows = grp.to_dict("records")
        # Mappe vorhandene Nummern pro Item-Typ in diesem Kapitel
        by_type = {}
        for idx, r in enumerate(grp_rows):
            iname = str(r.get("item_name") or "")
            m = ITEM_NUM_RE.match(iname)
            if m:
                raw_t = m.group(1).strip()
                norm_t = "Document" if "document" in raw_t.lower() else raw_t
                num = int(m.group(2))
                by_type.setdefault(norm_t, {})[num] = (
                    idx,
                    r.get("timestamp"),
                    iname,
                    r,
                )

        # Finde Lücken pro Typ
        inserts_after_idx = {}
        insert_at_start = []

        for itype, num_map in by_type.items():
            max_n = max(num_map.keys())
            if max_n < 2:
                continue
            missing = [n for n in range(1, max_n + 1) if n not in num_map]
            for m_n in missing:
                total_gaps += 1
                prev_info = num_map.get(m_n - 1)
                next_info = num_map.get(m_n + 1)
                prev_ts = prev_info[1] if prev_info else "00:00"
                next_ts = next_info[1] if next_info else "Kapitelende"
                template_row = (
                    (prev_info[3] if prev_info else next_info[3])
                    if (prev_info or next_info)
                    else grp_rows[0]
                )

                warn_row = dict(template_row)
                warn_row["item_name"] = (
                    f"{itype} #{m_n} [⚠️ FEHLT IM YOUTUBE-KAPITEL: Im Video zwischen {prev_ts} und {next_ts} prüfen!]"
                )
                warn_row["timestamp"] = prev_ts
                warn_row["category_group"] = infer_category_group(
                    f"{itype} #{m_n}", warn_row.get("category_group")
                )

                print(
                    f"  ⚠️ [{filename}] LÜCKE in '{chap}': {itype} #{m_n} fehlt (zwischen {prev_ts} und {next_ts}) -> Prüfzeile eingefügt!"
                )
                if prev_info:
                    inserts_after_idx.setdefault(prev_info[0], []).append(warn_row)
                else:
                    insert_at_start.append(warn_row)

        # Baue Kapitel-Zeilen inkl. eingefügter Warn-Zeilen zusammen
        for idx, r in enumerate(grp_rows):
            final_rows.append(r)
            if idx == 0 and insert_at_start:
                final_rows.extend(insert_at_start)
            if idx in inserts_after_idx:
                final_rows.extend(inserts_after_idx[idx])

    if total_gaps == 0:
        print(f"  ✅ [{filename}] Alle Kapitel-Sequenzen (#1..#N) sind lückenlos!")

    out_df = pd.DataFrame(final_rows)
    out_df["guide_id"] = range(1, len(out_df) + 1)
    return out_df


def style_output_excel(filepath: str):
    wb = openpyxl.load_workbook(filepath)
    header_fill = PatternFill(
        start_color="1E293B", end_color="1E293B", fill_type="solid"
    )
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    zebra_fill = PatternFill(
        start_color="F8FAFC", end_color="F8FAFC", fill_type="solid"
    )
    warn_fill = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="E2E8F0"),
        right=Side(style="thin", color="E2E8F0"),
        top=Side(style="thin", color="E2E8F0"),
        bottom=Side(style="thin", color="E2E8F0"),
    )

    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row > 1:
            ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center", horizontal="left")
        for row_idx in range(2, ws.max_row + 1):
            row_vals = [str(c.value or "") for c in ws[row_idx]]
            is_warn = any("FEHLT IM YOUTUBE" in v for v in row_vals)
            for cell in ws[row_idx]:
                cell.border = thin_border
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if is_warn:
                    cell.fill = warn_fill
                elif row_idx % 2 == 0:
                    cell.fill = zebra_fill

    col_widths = {
        "A": 15,
        "B": 18,
        "C": 10,
        "D": 14,
        "E": 36,
        "F": 34,
        "G": 45,
        "H": 68,
        "I": 12,
        "J": 28,
        "K": 10,
    }
    for sname in ("sheet_name=1", "sheet_name=2"):
        if sname in wb.sheetnames:
            ws = wb[sname]
            for col_l, w in col_widths.items():
                ws.column_dimensions[col_l].width = w

    wb.save(filepath)


def process_excel_file(filepath: str):
    base = os.path.basename(filepath)
    if base.startswith("~$") or any(
        base.endswith(suf) for suf in ("-clean.xlsx", "-v2.xlsx", "-v3.xlsx")
    ):
        return

    print(f"\n📂 Veredele: {base}...")
    xls = pd.ExcelFile(filepath)
    if "sheet_name=0" not in xls.sheet_names or "sheet_name=1" not in xls.sheet_names:
        print("  ⏭️ Übersprungen (keine sheet_name=0/1 Struktur).")
        return

    df0 = pd.read_excel(filepath, sheet_name="sheet_name=0")
    df1 = pd.read_excel(filepath, sheet_name="sheet_name=1")
    df3 = (
        pd.read_excel(filepath, sheet_name="sheet_name=3")
        if "sheet_name=3" in xls.sheet_names
        else pd.DataFrame()
    )

    # 1. Doppelte Dict-Strings in sheet_name=0 reparieren
    for col in ("icon_url", "global_seltenheit"):
        if col in df0.columns:
            df0[col] = (
                df0[col]
                .astype(str)
                .str.replace(
                    r"^\{\'de\':\s*\"(\{\'de\':\s*\'[^\']+\'\})\"\}$", r"\1", regex=True
                )
            )

    # 2. Kombi-Zeilen ('to', '-', '&', ',') splitten + Kapitel & Kategorien ableiten
    df1_expanded = expand_and_enrich_rows(df1)

    # 3. Kapitel-Sequenzen prüfen & fehlende Nummern als gelbe Warnzeile einfügen
    df1_final = detect_and_insert_gaps(df1_expanded, base)

    print(f"  📊 Zeilen in sheet_name=1: {len(df1)} roh -> {len(df1_final)} veredelt.")

    out_path = filepath[:-5] + "-clean.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as writer:
        df0.to_excel(writer, sheet_name="sheet_name=0", index=False)
        df1_final.to_excel(writer, sheet_name="sheet_name=1", index=False)
        df1_final.to_excel(writer, sheet_name="sheet_name=2", index=False)
        df3.to_excel(writer, sheet_name="sheet_name=3", index=False)

    style_output_excel(out_path)
    print(f"  💾 Gespeichert als: {os.path.basename(out_path)}")


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_TARGET_DIR
    if not os.path.exists(target):
        target = os.path.dirname(os.path.abspath(__file__))

    if os.path.isdir(target):
        excel_files = sorted(glob.glob(os.path.join(target, "*.xlsx")))
        print(f"🚀 Starte Ordner-Veredelung in: {target}")
        for f in excel_files:
            process_excel_file(f)
        print("\n🏁 Ordner-Durchlauf abgeschlossen!")
    else:
        process_excel_file(target)
